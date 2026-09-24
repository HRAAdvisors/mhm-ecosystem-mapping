"""
One-time export of the MHM Collaborations Tracker workbook into a static JSON
file consumed by the Next.js app (data/mhm-network.json).

Source (client-provided, not committed to this repo):
  OneDrive: Digital Opportunity - 260299 - MHM Broadband Infrastructure Mapping
            /3. Analysis/WS2/MHM Collaborations Tracker - Final.xlsx

Re-run this script and commit the regenerated JSON whenever the tracker changes:
  python3 scripts/export_tracker.py "/path/to/MHM Collaborations Tracker - Final.xlsx"
"""

import json
import re
import sys
from datetime import date
from pathlib import Path

import openpyxl

DEFAULT_SOURCE = (
    "/Users/eddiejoeantonio/Library/CloudStorage/OneDrive-SharedLibraries-HR&AAdvisors,Inc/"
    "Digital Opportunity - 260299 - MHM Broadband Infrastructure Mapping/"
    "3. Analysis/WS2/MHM Collaborations Tracker - Final.xlsx"
)

SHEET_NAME = "Master Grantee-Org Sheet"

# Maps each output field to the tracker's column header. Looked up by header
# name (not position) since the tracker's column order/count has changed
# before (e.g. the "Organization Region"/"Grantee Region" split).
#
# "region" is deliberately the ORGANIZATION's region, not the grantee's: the
# graph filters relationships by where the organization provides service
# (buildGraph in lib/data.ts), so a grantee headquartered in one region can
# still show a partner org whose own home region is elsewhere.
HEADER_MAP = {
    "grantee": "Grantee",
    "organization": "Organization",
    "organizationCounty": "Organization County",
    "region": "Organization Region",
    "regionSourceJustification": "Organization Location - Evidence/Confidence",
    "granteeRegion": "Grantee Region",
    "isGrantee": "Is Grantee?",
    "confirmationStatus": "Confirmation Status",
    "relationshipType": "Relationship Type",
    "relationshipTypeRationale": "Relationship Type - Rationale",
    "relationshipStrength": "Relationship Strength",
    "relationshipStrengthRationale": "Relationship Strength - Rationale",
    "additionalRegions": "Additional Region(s) (historical note)",
    "primaryServiceCategory": "Primary Service Category",
    "categoryConfidence": "Category Confidence",
    "categoryJustification": "Category - Justification/Evidence",
    "granteeFundingAmount": "Grantee Funding Amount",
    "fundingSource": "Funding - Source",
    "evidenceQuote": "Evidence / Quote",
    "sourceDocuments": "Source Document(s)",
    "locationCited": "Location Cited in Source",
    "newVsExisting": "New vs. Existing",
    "notesFlags": "Notes / Flags",
    "activeGrant2026": "Active Grant (2026)?",
    "otherMhmGranteeStatus": "Other MHM Grantee Status",
}

REGION_RE = re.compile(r"^Region ([A-L])\b")

REGION_LABELS = {
    "A": "Region A — Concho Valley / San Angelo",
    "B": "Region B — Texas Hill Country",
    "C": "Region C — Travis County / Austin",
    "D": "Region D — Bexar County / San Antonio",
    "E": "Region E — Bastrop-Hays-Fayette",
    "F": "Region F — Medina-Atascosa-Guadalupe",
    "G": "Region G — Victoria / Coastal Plains",
    "H": "Region H — Corpus Christi / Coastal Bend",
    "J": "Region J — South Texas / Rio Grande Valley",
    "K": "Region K — Tri-County (Laredo)",
    "L": "Region L — Mid-Border Region",
}


def region_code(raw):
    if not raw:
        return None
    m = REGION_RE.match(str(raw).strip())
    return m.group(1) if m else None


def additional_region_codes(raw):
    if not raw:
        return []
    codes = []
    for part in str(raw).split(";"):
        c = region_code(part.strip())
        if c and c not in codes:
            codes.append(c)
    return codes


# Matches "Region A/B", "Region A / B", "Region A, B", and also the bare
# form some rows use with no "Region" prefix at all ("A/B") -- a grantee that
# genuinely serves more than one region written as one phrase (as opposed to
# `additional_region_codes`, which handles a semicolon-separated list of
# several *separate* "Region X" labels).
MULTI_REGION_RE = re.compile(r"^(?:Region\s+)?([A-L](?:\s*[/,]\s*[A-L])+)\b")


def grantee_region_codes(raw):
    """The grantee's own home region(s): usually exactly one, but a cell like
    "Region A/B - ..." or the bare "A/B" means the grantee serves both.
    Returns (primary_code, additional_codes)."""
    if not raw:
        return None, []
    text = str(raw).strip()
    m = MULTI_REGION_RE.match(text)
    if m:
        codes = [c.strip() for c in re.split(r"[/,]", m.group(1))]
        return codes[0], codes[1:]
    return region_code(text), []


def clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def find_section_bounds(ws):
    """Locate the single header row ('Grantee' in col A) and the row range for
    the relationship section, then the "KEY REGIONAL PLAYERS" title row that
    starts the second section (blank Grantee column from there on)."""
    header_row = None
    section_title_row = None
    for r in range(1, ws.max_row + 1):
        value = ws.cell(row=r, column=1).value
        if value == "Grantee":
            header_row = r
        elif value and "KEY REGIONAL PLAYERS" in str(value).upper():
            section_title_row = r
            break
    if header_row is None:
        raise RuntimeError("Could not find the 'Grantee' header row")
    if section_title_row is None:
        raise RuntimeError("Could not find the 'KEY REGIONAL PLAYERS' section title row")

    relationship_start = header_row + 1
    relationship_end = section_title_row - 1
    krp_start = section_title_row + 1
    krp_end = ws.max_row
    return (relationship_start, relationship_end), (krp_start, krp_end)


def build_column_index(ws, header_row):
    """Maps each output field to its column number by matching HEADER_MAP
    against the actual header row text, so a reordered/inserted column in the
    tracker doesn't silently shift every field over (as happened when
    "Organization Region"/"Grantee Region" were split out)."""
    header_to_col = {}
    for c in range(1, ws.max_column + 1):
        value = ws.cell(row=header_row, column=c).value
        if value:
            header_to_col[str(value).strip()] = c

    column_index = {}
    missing = []
    for field, header in HEADER_MAP.items():
        col = header_to_col.get(header)
        if col is None:
            missing.append(header)
        else:
            column_index[field] = col
    if missing:
        raise RuntimeError(f"Tracker headers not found: {missing}")
    return column_index


def extract_rows(ws, start, end, section, column_index):
    rows = []
    for r in range(start, end + 1):
        values = {field: clean(ws.cell(row=r, column=c).value) for field, c in column_index.items()}
        if all(v is None for v in values.values()):
            continue
        # Keep a row if it has a partner organization (the normal case), or
        # if it's a grantee's own summary row with no documented partner yet
        # (organization blank, grantee filled) -- otherwise that grantee's
        # own category/region can never be displayed. See lookupSubsector in
        # lib/data.ts, which reads such rows for a grantee's own category.
        if not values.get("organization") and not values.get("grantee"):
            continue
        row = values
        row["section"] = section
        row["sourceRow"] = r
        # The "region"/"additionalRegions" fields describe the ORGANIZATION's
        # location and are meaningless without one -- a stray value there on
        # an org-less grantee self-row would otherwise make
        # organizationServesRegion() match on a null organization (lib/data.ts),
        # inserting `null` into a region's node set and crashing name.localeCompare.
        if row.get("organization"):
            row["regionCode"] = region_code(row.get("region"))
            row["additionalRegionCodes"] = additional_region_codes(row.get("additionalRegions"))
        else:
            row["regionCode"] = None
            row["additionalRegionCodes"] = []
        row["granteeRegionCode"], row["granteeAdditionalRegionCodes"] = grantee_region_codes(
            row.get("granteeRegion")
        )
        rows.append(row)
    return rows


def main():
    source = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_SOURCE
    out_path = Path(__file__).resolve().parent.parent / "data" / "mhm-network.json"

    wb = openpyxl.load_workbook(source, data_only=True)
    ws = wb[SHEET_NAME]

    (rel_start, rel_end), (krp_start, krp_end) = find_section_bounds(ws)
    column_index = build_column_index(ws, rel_start - 1)
    rows = extract_rows(ws, rel_start, rel_end, "relationship", column_index)
    rows += extract_rows(ws, krp_start, krp_end, "key_regional_player", column_index)

    present_codes = set()
    for row in rows:
        if row["regionCode"]:
            present_codes.add(row["regionCode"])
        present_codes.update(row["additionalRegionCodes"])

    regions = [
        {"code": code, "label": REGION_LABELS.get(code, f"Region {code}")}
        for code in sorted(present_codes)
    ]

    categories = sorted(
        {
            row["primaryServiceCategory"]
            for row in rows
            if row["primaryServiceCategory"]
            and "N/A" not in row["primaryServiceCategory"]
            and row["primaryServiceCategory"] != "Unable to classify"
        }
    )

    payload = {
        "generatedAt": date.today().isoformat(),
        "sourceSheet": SHEET_NAME,
        "regions": regions,
        "categories": categories,
        "rows": rows,
    }

    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False))
    print(f"Wrote {len(rows)} rows ({len(regions)} regions) to {out_path}")


if __name__ == "__main__":
    main()
